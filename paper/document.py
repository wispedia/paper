import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn
from latex2mathml.converter import convert as latex_to_mathml
from mathml2omml import convert as mathml_to_omml
from PIL import Image


MATH = re.compile(r"(\$\$.*?\$\$|\$[^$]+\$)", re.S)


def inline(paragraph, text):
    for part in MATH.split(text):
        if not part:
            continue
        if part.startswith("$") and part.endswith("$"):
            latex = part.strip("$").strip()
            try:
                value = mathml_to_omml(latex_to_mathml(latex))
                # mathml2omml 0.0.2 closes groupChrPr as groupChr for overbars.
                # Repair only that malformed properties fragment, not valid groups.
                value = re.sub(r"(<m:groupChrPr>(?:<m:\w+\b[^>]*/>)*)</m:groupChr>", r"\1</m:groupChrPr>", value)
                # mathml2omml returns the native Word equation element.
                xml = '<root xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">' + value + '</root>'
                root = parse_xml(xml)
                # OMML radicals require a degree element, including square roots.
                # mathml2omml omits it for msqrt; readers then misparse nested fractions.
                for radical in root.findall(".//" + qn("m:rad")):
                    if radical.find(qn("m:deg")) is None:
                        properties = radical.find(qn("m:radPr"))
                        if properties is None:
                            properties = OxmlElement("m:radPr")
                            radical.insert(0, properties)
                        hidden = OxmlElement("m:degHide")
                        hidden.set(qn("m:val"), "1")
                        properties.append(hidden)
                        radical.insert(list(radical).index(properties) + 1, OxmlElement("m:deg"))
                for child in root:
                    paragraph._p.append(child)
            except Exception as exc:
                raise ValueError("部分公式无法排版，请重试或更换清晰图片。") from exc
        else:
            paragraph.add_run(part)


def new_document(content, identity=True):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(8)
    for name, size in [("Title", 19), ("Heading 1", 13), ("Heading 2", 11)]:
        style = doc.styles[name]
        style.font.name = "SimSun"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    p = doc.add_paragraph(content["title"], style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if content["subtitle"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        inline(p, content["subtitle"])
    if identity:
        doc.add_paragraph("班级：____________　姓名：____________　学号：____________")
    return doc


def build_document(content, directory):
    directory = Path(directory)
    doc = new_document(content)
    for sec in content["sections"]:
        if sec["heading"]:
            doc.add_paragraph(sec["heading"], style="Heading 1")
        if sec["instructions"]:
            inline(doc.add_paragraph(), sec["instructions"])
        for question in sec["questions"]:
            p = doc.add_paragraph()
            p.paragraph_format.keep_together = True
            p.add_run((question["number"] + ". ") if question["number"] else "")
            inline(p, question["stem"])
            if question["options"]:
                p.paragraph_format.keep_with_next = True
                compact = all(len(option["text"]) <= 65 and len(MATH.sub("", option["text"])) <= 12 for option in question["options"])
                step = 2 if compact else 1
                for start in range(0, len(question["options"]), step):
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(.5)
                    p.paragraph_format.keep_together = True
                    p.paragraph_format.keep_with_next = start + step < len(question["options"])
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(8))
                    for index, option in enumerate(question["options"][start:start+step]):
                        if index:
                            p.add_run("\t")
                        p.add_run(option["label"] + ". ")
                        inline(p, option["text"])
            for figure_index, figure in enumerate(question["figures"]):
                x1, y1, x2, y2 = figure["bbox"]
                if not (0 <= x1 < x2 <= 1000 and 0 <= y1 < y2 <= 1000):
                    raise ValueError("试卷配图的位置需要重新识别，请重试。")
                with Image.open(directory / "input.jpg") as original:
                    w, h = original.size
                    crop = original.crop((x1*w/1000, y1*h/1000, x2*w/1000, y2*h/1000))
                    # Question numbers come from model output; never use them as paths.
                    target = directory / f"figure-{len(doc.inline_shapes)}.png"
                    crop.save(target)
                    width = min(12, max(4, crop.width / w * 17))
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(target), width=Cm(width))
                if figure["caption"]:
                    inline(doc.add_paragraph(), figure["caption"])
            for subquestion in question["subquestions"]:
                p = doc.add_paragraph()
                p.paragraph_format.keep_together = True
                inline(p, subquestion)
                doc.add_paragraph("\n\n")
    if content["uncertainties"]:
        doc.add_paragraph("待核对内容", style="Heading 1")
        for note in content["uncertainties"]:
            doc.add_paragraph(note)
    doc.core_properties.title = content["title"]
    doc.core_properties.author = ""
    path = directory / "paper.docx"
    doc.save(path)
    return path


def handwriting_paragraphs(text):
    """Reflow prose outside equations; never split a multiline math construct."""
    lines, current = [], ""
    for token in MATH.split(text):
        if token.startswith("$") and token.endswith("$"):
            current += token
            continue
        fragments = token.split("\n")
        current += fragments[0]
        for fragment in fragments[1:]:
            lines.append(current.strip())
            current = fragment
    lines.append(current.strip())
    blocks = []
    prose = ""
    for line in lines:
        # Pure calculation lines remain steps; natural-language lines wrap normally.
        formula = bool(MATH.search(line)) and not MATH.sub("", line).strip(" 。；，;,. ")
        if not line or formula:
            if prose:
                blocks.append(prose)
                prose = ""
            if line:
                blocks.append(line)
        else:
            # A line break between Latin words needs a separator; Chinese does not.
            separator = " " if prose and prose[-1:].isascii() and line[:1].isascii() else ""
            prose += separator + line
    if prose:
        blocks.append(prose)
    return blocks


def build_handwriting_document(content, directory):
    from .handwriting import warnings

    doc = new_document(content, identity=False)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.15
    normal.space_after = Pt(4)
    title = doc.styles["Title"]
    title.font.size = Pt(17)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    heading = doc.styles["Heading 1"].paragraph_format
    heading.space_before = Pt(11)
    heading.space_after = Pt(5)
    heading.keep_with_next = True
    for answer in content["answers"]:
        if answer["question"]:
            doc.add_paragraph("第" + answer["question"] + "题", style="Heading 1")
        for part in answer["parts"]:
            blocks = handwriting_paragraphs(part["text"])
            for index, block in enumerate(blocks or [""]):
                p = doc.add_paragraph()
                fmt = p.paragraph_format
                fmt.left_indent = Cm(.85)
                if index == 0 and part["label"]:
                    fmt.space_before = Pt(3)
                    fmt.first_line_indent = Cm(-.85)
                    fmt.tab_stops.add_tab_stop(Cm(.85))
                    p.add_run(part["label"] + "\t")
                inline(p, block)
                fmt.keep_together = bool(MATH.search(block)) and not MATH.sub("", block).strip(" 。；，;,. ")
        # Legacy results may contain crossed_out; it is deliberately never exported.
    notes = warnings(content)
    if notes:
        doc.add_paragraph("待核对内容", style="Heading 1")
        for note in notes:
            p = doc.add_paragraph(note)
            p.paragraph_format.space_after = Pt(3)
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(90, 90, 90)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = content["title"]
    doc.core_properties.author = ""
    path = Path(directory) / "paper.docx"
    doc.save(path)
    return path
